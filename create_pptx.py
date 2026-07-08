from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()

# Slide 1: Title Slide
slide_layout = prs.slide_layouts[0] # Title slide
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "FinPersona: Credit Risk Prediction Module"
subtitle.text = "Project Discussion & Evaluation"

# Slide 2: Data Pipeline & Feature Engineering
slide_layout = prs.slide_layouts[1] # Title and Content
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
title.text = "Data Pipeline & Feature Engineering"
content = slide.placeholders[1]
tf = content.text_frame
tf.text = "Foundation: Merging Prime Profiles & Transactions"
p = tf.add_paragraph()
p.text = "Ingesting monthly credit transactions"
p.level = 1
p = tf.add_paragraph()
p.text = "Merging with demographic 'prime' profiles"
p.level = 1
p = tf.add_paragraph()
p.text = "Feature Engineering"
p.level = 0
p = tf.add_paragraph()
p.text = "Extracting patterns: utilization rates, payment histories"
p.level = 1
p = tf.add_paragraph()
p.text = "Robust Data Cleaning Pipeline"
p.level = 0
p = tf.add_paragraph()
p.text = "Handling anomalies and normalizing features across time periods"
p.level = 1

# Slide 3: Model Architecture & UI Integration
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
title.text = "Model Architecture & UI Integration"
content = slide.placeholders[1]
tf = content.text_frame
tf.text = "Handling Imbalanced Data"
p = tf.add_paragraph()
p.text = "Defaults are rare: specialized sampling techniques used"
p.level = 1
p = tf.add_paragraph()
p.text = "Allows model to learn risky profiles without majority class bias"
p.level = 1
p = tf.add_paragraph()
p.text = "UI Integration"
p.level = 0
p = tf.add_paragraph()
p.text = "Modeling pipeline wrapped in a Flask API"
p.level = 1
p = tf.add_paragraph()
p.text = "Real-time predictions served directly to FinPersona web interface"
p.level = 1

# Slide 4: Evaluation & Performance Results
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
title.text = "Evaluation & Performance Results"
content = slide.placeholders[1]
tf = content.text_frame
tf.text = "Overall Accuracy: 96%"
p = tf.add_paragraph()
p.text = "Non-Default: Precision 0.99 | Recall 0.97 | F1 0.98"
p.level = 1
p = tf.add_paragraph()
p.text = "Default: Precision 0.77 | Recall 0.84 | F1 0.80"
p.level = 1
p = tf.add_paragraph()
p.text = "The F-beta Score Advantage"
p.level = 0
p = tf.add_paragraph()
p.text = "Weighs recall higher than precision"
p.level = 1
p = tf.add_paragraph()
p.text = "Crucial for risk management: False negatives are highly expensive"
p.level = 1
p = tf.add_paragraph()
p.text = "Optimizing F-beta secures an 84% default capture rate"
p.level = 1

# Slide 5: Conclusion & Next Steps
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
title.text = "Conclusion & Next Steps"
content = slide.placeholders[1]
tf = content.text_frame
tf.text = "Summary"
p = tf.add_paragraph()
p.text = "End-to-end credit risk pipeline managing imbalanced data"
p.level = 1
p = tf.add_paragraph()
p.text = "Seamless integration into the FinPersona UI"
p.level = 1
p = tf.add_paragraph()
p.text = "Successfully identifies 84% of potential defaults"
p.level = 1
p = tf.add_paragraph()
p.text = "Future Work"
p.level = 0
p = tf.add_paragraph()
p.text = "Incorporate granular transaction features to further boost recall"
p.level = 1

prs.save("Credit_Risk_Presentation.pptx")
print("Presentation generated successfully at Credit_Risk_Presentation.pptx")
